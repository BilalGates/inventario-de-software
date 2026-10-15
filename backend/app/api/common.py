"""Utilidades comunes de la API: paginación, ordenación y exportación."""
from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from typing import Any

from fastapi import HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import Select, func, select
from sqlalchemy.orm import Session


class Page:
    def __init__(self, page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=1000),
                 sort: str | None = Query(None, description="campo o -campo")):
        self.page, self.page_size, self.sort = page, page_size, sort


def paginate(db: Session, query: Select, page: Page, sortable: dict[str, Any], default_sort: Any) -> tuple[list, int]:
    total = db.execute(select(func.count()).select_from(query.order_by(None).subquery())).scalar() or 0
    order = default_sort
    if page.sort:
        key = page.sort.lstrip("-")
        if key not in sortable:
            raise HTTPException(400, f"No se puede ordenar por '{key}'")
        order = sortable[key].desc() if page.sort.startswith("-") else sortable[key].asc()
    rows = db.execute(query.order_by(order).offset((page.page - 1) * page.page_size).limit(page.page_size)).all()
    return rows, total


def not_found(what: str = "Recurso"):
    return HTTPException(404, f"{what} no encontrado")


def export_response(rows: Iterable[dict], columns: list[tuple[str, str]], filename: str, fmt: str) -> StreamingResponse:
    """columns: [(clave, título)]. fmt: csv | xlsx."""
    rows = list(rows)
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos"
        ws.append([title for _, title in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F2937")
        for row in rows:
            ws.append([_csv_safe(_cell(row.get(k))) for k, _ in columns])  # openpyxl trata "=..." como fórmula
        for i, (_, title) in enumerate(columns, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(60, len(title) + 4))
        ws.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'})
    buf = io.StringIO()
    buf.write("﻿")  # Excel abre bien el UTF-8
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([title for _, title in columns])
    for row in rows:
        writer.writerow([_csv_safe(_cell(row.get(k))) for k, _ in columns])
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": f'attachment; filename="{filename}.csv"'})


def _cell(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return str(value)
    return value


def _csv_safe(value: Any) -> Any:
    """Evita inyección de fórmulas al abrir el CSV en Excel."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value
