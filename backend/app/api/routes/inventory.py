"""Ejecuciones, fuentes, inventarios mensuales, cambios, alertas, casos de revisión y auditoría."""
from __future__ import annotations

import io
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from app.api.common import Page, _csv_safe, export_response, not_found, paginate
from app.connectors.files import decode
from app.core.config import get_settings
from app.core.security import Admin, CurrentUser, Reader, Technician
from app.db import get_db, utcnow
from app.models import (
    Alert,
    AppSetting,
    AuditEvent,
    Department,
    Device,
    DeviceAssignment,
    DeviceExternalId,
    InventoryChange,
    InventoryRun,
    InventorySnapshot,
    InventorySnapshotItem,
    InventorySource,
    ReviewCase,
    SoftwareInstallation,
    SoftwareProduct,
)
from app.services import audit, snapshots
from app.services import runs as run_service
from app.services.bootstrap import get_setting

router = APIRouter(prefix="/api", tags=["Inventario"])
MAX_UPLOAD = 20 * 1024 * 1024


# --------------------------------------------------------------------------- ejecuciones
def _run(r: InventoryRun, full: bool = False) -> dict:
    data = {"id": r.id, "source": r.source.code, "source_name": r.source.name, "trigger": r.trigger,
            "status": r.status, "is_complete": r.is_complete, "requested_by": r.requested_by,
            "requested_at": r.requested_at, "started_at": r.started_at, "finished_at": r.finished_at,
            "attempt": r.attempt, "retry_of_id": r.retry_of_id, "period": r.period, "error": r.error,
            "stats": r.stats or {}, "filename": (r.params or {}).get("filename")}
    if full:
        data["log"] = r.log or []
        data["params"] = {k: v for k, v in (r.params or {}).items() if k != "content"}
    return data


@router.get("/runs")
def list_runs(page: Page = Depends(), source: str | None = None, status: str | None = None,
              db: Session = Depends(get_db), _: CurrentUser = Reader):
    query = select(InventoryRun)
    if source:
        query = query.join(InventorySource).where(InventorySource.code == source)
    if status:
        query = query.where(InventoryRun.status.in_(status.split(",")))
    rows, total = paginate(db, query, page, {"requested_at": InventoryRun.requested_at}, InventoryRun.id.desc())
    return {"total": total, "items": [_run(r[0]) for r in rows]}


@router.get("/runs/{run_id}")
def get_run(run_id: int, db: Session = Depends(get_db), _: CurrentUser = Reader):
    run = db.get(InventoryRun, run_id)
    if run is None:
        raise not_found("Ejecución")
    changes = dict(db.execute(select(InventoryChange.change_type, func.count()).where(InventoryChange.run_id == run_id)
                              .group_by(InventoryChange.change_type)).all())
    return {**_run(run, full=True), "changes": changes}


class SyncIn(BaseModel):
    source: str = Field(..., pattern="^[a-z_]+$")


@router.post("/runs", status_code=202)
def start_sync(body: SyncIn, db: Session = Depends(get_db), user: CurrentUser = Technician):
    """Encola una sincronización manual; la ejecuta el worker."""
    src = db.execute(select(InventorySource).where(InventorySource.code == body.source)).scalar_one_or_none()
    if src is None:
        raise not_found("Fuente")
    if src.kind != "api":
        raise HTTPException(422, f"{src.name} se alimenta con ficheros: use la subida de ficheros")
    if src.code == "intune" and not get_settings().graph_configured:
        raise HTTPException(409, "El conector de Intune no tiene credenciales configuradas (ver docs/integrations.md)")
    run = run_service.enqueue(db, src.code, "manual", user.upn)
    db.commit()
    return _run(run)


@router.post("/runs/upload", status_code=202)
async def upload_file(source: str = Form(..., pattern="^(panda|intune)$"), file: UploadFile = File(...),
                      device_id: int | None = Form(None), declared_complete: bool = Form(False),
                      db: Session = Depends(get_db), user: CurrentUser = Technician):
    """Ingesta de ficheros: exportación de Panda (un equipo) o CSV de Intune (dispositivos o aplicaciones)."""
    raw = await file.read(MAX_UPLOAD + 1)
    if len(raw) > MAX_UPLOAD:
        raise HTTPException(413, "Fichero demasiado grande (máx. 20 MB)")
    try:
        content = decode(raw)
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc
    if device_id is not None and db.get(Device, device_id) is None:
        raise not_found("Dispositivo")
    params = {"content": content, "filename": file.filename, "declared_complete": declared_complete}
    if device_id:
        params["device_id"] = device_id
    run = run_service.enqueue(db, source, "upload", user.upn, params)
    db.commit()
    return _run(run)


@router.post("/runs/{run_id}/retry", status_code=202)
def retry_run(run_id: int, db: Session = Depends(get_db), user: CurrentUser = Technician):
    try:
        run = run_service.retry(db, run_id, user.upn)
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from exc
    db.commit()
    return _run(run)


# --------------------------------------------------------------------------- fuentes
@router.get("/sources")
def list_sources(db: Session = Depends(get_db), _: CurrentUser = Reader):
    s = get_settings()
    out = []
    for src in db.execute(select(InventorySource).order_by(InventorySource.priority)).scalars():
        last = db.execute(select(InventoryRun).where(InventoryRun.source_id == src.id)
                          .order_by(InventoryRun.id.desc()).limit(1)).scalar_one_or_none()
        # Fuentes sin identificadores oficiales (Panda) solo se reflejan a través de su software.
        devices = len(set(db.execute(select(DeviceExternalId.device_id).where(DeviceExternalId.source == src.code)).scalars())
                      | set(db.execute(select(SoftwareInstallation.device_id)
                                       .where(SoftwareInstallation.source == src.code)).scalars()))
        devices_sw = db.execute(select(func.count(func.distinct(SoftwareInstallation.device_id)))
                                .where(SoftwareInstallation.source == src.code,
                                       SoftwareInstallation.present.is_(True))).scalar() or 0
        configured = s.graph_configured if src.code == "intune" else True
        out.append({"id": src.id, "code": src.code, "name": src.name, "kind": src.kind, "enabled": src.enabled,
                    "priority": src.priority, "config": src.config or {}, "configured": configured,
                    "last_success_at": src.last_success_at, "last_error": src.last_error,
                    "last_run": _run(last) if last else None, "devices": devices,
                    "devices_with_software": devices_sw,
                    "watch_folder": bool(s.panda_drop_dir) if src.code == "panda" else None})
    return out


class SourcePatch(BaseModel):
    enabled: bool | None = None
    priority: int | None = Field(None, ge=1, le=1000)
    config: dict | None = None


@router.patch("/sources/{code}")
def update_source(code: str, body: SourcePatch, db: Session = Depends(get_db), user: CurrentUser = Admin):
    src = db.execute(select(InventorySource).where(InventorySource.code == code)).scalar_one_or_none()
    if src is None:
        raise not_found("Fuente")
    if body.config is not None:
        allowed = {"software_strategy", "api_version", "fetch_hardware_details"}
        if set(body.config) - allowed:
            raise HTTPException(422, f"Claves permitidas: {sorted(allowed)}. Los secretos van en variables de entorno.")
        if body.config.get("software_strategy", "per_app") not in ("per_app", "per_device"):
            raise HTTPException(422, "software_strategy debe ser per_app o per_device")
        if body.config.get("api_version", "v1.0") not in ("v1.0", "beta"):
            raise HTTPException(422, "api_version debe ser v1.0 o beta")
    before = audit.snapshot(src, ["enabled", "priority", "config"])
    for key, value in body.model_dump(exclude_unset=True).items():
        setattr(src, key, {**(src.config or {}), **value} if key == "config" else value)
    audit.record(db, user.upn, "source.update", "inventory_source", src.id, src.name, before=before,
                 after=audit.snapshot(src, ["enabled", "priority", "config"]))
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------- configuración
SETTINGS_SCHEMA = {"schedule.monthly": dict, "alerts.stale_source_days": int}


@router.get("/settings")
def get_settings_api(db: Session = Depends(get_db), _: CurrentUser = Reader):
    out = {key: get_setting(db, key) for key in SETTINGS_SCHEMA}
    hb = db.get(AppSetting, "worker.heartbeat")
    out["worker.heartbeat"] = hb.value if hb else None
    s = get_settings()
    out["runtime"] = {"environment": s.environment, "auth_mode": s.auth_mode, "graph_configured": s.graph_configured,
                      "panda_drop_dir_configured": bool(s.panda_drop_dir), "stale_device_days": s.stale_device_days}
    return out


class ScheduleIn(BaseModel):
    enabled: bool
    day: int = Field(..., ge=1, le=28)
    hour: int = Field(..., ge=0, le=23)
    sources: list[str]


@router.put("/settings/schedule")
def set_schedule(body: ScheduleIn, db: Session = Depends(get_db), user: CurrentUser = Admin):
    valid = set(db.execute(select(InventorySource.code).where(InventorySource.kind == "api")).scalars())
    if set(body.sources) - valid:
        raise HTTPException(422, f"Fuentes programables: {sorted(valid)}")
    row = db.get(AppSetting, "schedule.monthly") or AppSetting(key="schedule.monthly")
    before = row.value
    row.value, row.updated_by = body.model_dump(), user.upn
    db.add(row)
    audit.record(db, user.upn, "settings.schedule", "app_setting", "schedule.monthly", "Programación mensual",
                 before={"value": before}, after={"value": row.value})
    db.commit()
    return row.value


@router.put("/settings/stale-days")
def set_stale(days: int = Query(..., ge=1, le=60), db: Session = Depends(get_db), user: CurrentUser = Admin):
    row = db.get(AppSetting, "alerts.stale_source_days") or AppSetting(key="alerts.stale_source_days")
    row.value, row.updated_by = days, user.upn
    db.add(row)
    audit.record(db, user.upn, "settings.update", "app_setting", row.key, f"stale_source_days={days}")
    db.commit()
    return {"ok": True}


# --------------------------------------------------------------------------- inventarios mensuales
def _snap(s: InventorySnapshot) -> dict:
    return {"id": s.id, "period": s.period, "version": s.version, "status": s.status, "is_complete": s.is_complete,
            "coverage": s.coverage, "totals": s.totals, "created_at": s.created_at, "created_by": s.created_by,
            "reviewed_by": s.reviewed_by, "reviewed_at": s.reviewed_at, "review_notes": s.review_notes}


@router.get("/snapshots")
def list_snapshots(db: Session = Depends(get_db), _: CurrentUser = Reader):
    return [_snap(s) for s in db.execute(select(InventorySnapshot).order_by(
        InventorySnapshot.period.desc(), InventorySnapshot.version.desc())).scalars()]


@router.get("/snapshots/coverage")
def snapshot_coverage(period: str = Query(..., pattern=r"^\d{4}-(0[1-9]|1[0-2])$"), db: Session = Depends(get_db),
                      _: CurrentUser = Reader):
    return snapshots.coverage(db, period)


class SnapshotIn(BaseModel):
    period: str = Field(..., pattern=r"^\d{4}-(0[1-9]|1[0-2])$")


@router.post("/snapshots", status_code=201)
def create_snapshot(body: SnapshotIn, db: Session = Depends(get_db), user: CurrentUser = Technician):
    snap = snapshots.generate(db, body.period, user.upn)
    db.commit()
    return _snap(snap)


class ReviewIn(BaseModel):
    approve: bool
    notes: str | None = Field(None, max_length=4000)
    accept_incomplete: bool = False


@router.post("/snapshots/{snapshot_id}/review")
def review_snapshot(snapshot_id: int, body: ReviewIn, db: Session = Depends(get_db), user: CurrentUser = Technician):
    try:
        snap = snapshots.review(db, snapshot_id, user.upn, body.approve, body.notes, body.accept_incomplete)
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from exc
    db.commit()
    return _snap(snap)


def _names(db: Session, device_ids: set[int], product_ids: set[int]) -> tuple[dict, dict]:
    devs = {d.id: d.display_name for d in db.execute(select(Device).where(Device.id.in_(device_ids))).scalars()} \
        if device_ids else {}
    prods = dict(db.execute(select(SoftwareProduct.id, SoftwareProduct.name)
                            .where(SoftwareProduct.id.in_(product_ids))).all()) if product_ids else {}
    return devs, prods


@router.get("/snapshots/compare")
def compare_snapshots(base: int, target: int, db: Session = Depends(get_db), _: CurrentUser = Reader):
    if not db.get(InventorySnapshot, base) or not db.get(InventorySnapshot, target):
        raise not_found("Inventario")
    diff = snapshots.compare(db, base, target)
    dev_ids = set(diff["devices_added"]) | set(diff["devices_removed"]) | {
        x["device_id"] for k in ("installs_added", "installs_removed", "versions_changed") for x in diff[k]}
    prod_ids = set(diff["products_new"]) | set(diff["products_gone"]) | {
        x["product_id"] for k in ("installs_added", "installs_removed", "versions_changed") for x in diff[k]}
    devs, prods = _names(db, dev_ids, prod_ids)
    for k in ("installs_added", "installs_removed", "versions_changed"):
        for x in diff[k]:
            x["device"], x["product"] = devs.get(x["device_id"]), prods.get(x["product_id"])
    diff["devices_added"] = [{"id": i, "name": devs.get(i)} for i in diff["devices_added"]]
    diff["devices_removed"] = [{"id": i, "name": devs.get(i)} for i in diff["devices_removed"]]
    diff["products_new"] = [{"id": i, "name": prods.get(i)} for i in diff["products_new"]]
    diff["products_gone"] = [{"id": i, "name": prods.get(i)} for i in diff["products_gone"]]
    return diff


@router.get("/snapshots/{snapshot_id}/report")
def snapshot_report(snapshot_id: int, db: Session = Depends(get_db), user: CurrentUser = Reader):
    """Informe mensual en Excel: resumen, cobertura, equipos, software, cambios y no autorizados."""
    snap = db.get(InventorySnapshot, snapshot_id)
    if snap is None:
        raise not_found("Inventario")
    items = db.execute(select(InventorySnapshotItem).where(InventorySnapshotItem.snapshot_id == snapshot_id)).scalars().all()
    dev_ids = {i.device_id for i in items}
    prod_ids = {i.product_id for i in items if i.product_id}
    devices = {d.id: d for d in db.execute(select(Device).where(Device.id.in_(dev_ids))).scalars()} if dev_ids else {}
    products = {p.id: p for p in db.execute(select(SoftwareProduct).where(SoftwareProduct.id.in_(prod_ids))).scalars()} \
        if prod_ids else {}
    depts = dict(db.execute(select(Department.id, Department.name)).all())

    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")

    def sheet(title, headers, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append([_csv_safe(v) for v in r])
        ws.freeze_panes = "A2"
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 28
        return ws

    t = snap.totals or {}
    sheet("Resumen", ["Concepto", "Valor"], [
        ("Periodo", snap.period), ("Versión", snap.version), ("Estado", snap.status),
        ("Completo", "Sí" if snap.is_complete else "NO — alguna fuente esencial no completó"),
        ("Generado", f"{snap.created_at:%Y-%m-%d %H:%M} UTC por {snap.created_by}"),
        ("Revisado", f"{snap.reviewed_at:%Y-%m-%d %H:%M} UTC por {snap.reviewed_by}" if snap.reviewed_at else "Pendiente"),
        ("Notas de revisión", snap.review_notes or ""), ("Dispositivos", t.get("devices")),
        ("Instalaciones", t.get("installations")), ("Productos", t.get("products")),
        *[(f"Instalaciones {k}", v) for k, v in sorted((t.get("by_status") or {}).items())],
        ("Informe generado", f"{utcnow():%Y-%m-%d %H:%M} UTC por {user.upn}"),
        ("Nota", "Evidencia del inventario; no constituye por sí misma una declaración de conformidad ENS."),
    ], first=True)
    sheet("Cobertura", ["Fuente", "Esencial", "Correcta", "Ejecuciones", "Última completa", "Fallos posteriores"], [
        (v["name"], v["essential"], v["ok"], v["runs"], v["last_complete_run_id"], v["failed_after_last_complete"])
        for v in ((snap.coverage or {}).get("sources") or {}).values()])
    sheet("Dispositivos", ["Equipo", "Nº serie", "Departamento", "Usuario", "Modelo", "SO", "Última observación"], [
        (d.display_name, d.serial_number, depts.get(i.department_id), d.primary_user.upn if d.primary_user else None,
         d.model, f"{d.os_name or ''} {d.os_version or ''}".strip(), d.last_seen_at)
        for i in items if i.product_id is None and (d := devices.get(i.device_id))])
    sheet("Software", ["Equipo", "Departamento", "Producto", "Editor", "Versión", "Autorización"], [
        (devices[i.device_id].display_name, depts.get(i.department_id), products[i.product_id].name,
         products[i.product_id].publisher, i.version, i.authorization_status)
        for i in items if i.product_id and products[i.product_id].is_relevant])
    sheet("No autorizado", ["Equipo", "Producto", "Versión"], [
        (devices[i.device_id].display_name, products[i.product_id].name, i.version)
        for i in items if i.product_id and i.authorization_status == "unauthorized"])
    prev = snapshots.previous_snapshot(db, snap)
    if prev:
        diff = snapshots.compare(db, prev.id, snap.id)
        all_dev = {x["device_id"] for k in ("installs_added", "installs_removed", "versions_changed") for x in diff[k]}
        all_prod = {x["product_id"] for k in ("installs_added", "installs_removed", "versions_changed") for x in diff[k]}
        dn, pn = _names(db, all_dev, all_prod)
        rows = [("Alta", dn.get(x["device_id"]), pn.get(x["product_id"]), "", ", ".join(x["versions"]))
                for x in diff["installs_added"]]
        rows += [("Baja", dn.get(x["device_id"]), pn.get(x["product_id"]), ", ".join(x["versions"]), "")
                 for x in diff["installs_removed"]]
        rows += [("Versión", dn.get(x["device_id"]), pn.get(x["product_id"]), ", ".join(x["from"]), ", ".join(x["to"]))
                 for x in diff["versions_changed"]]
        sheet(f"Cambios vs {prev.period}", ["Tipo", "Equipo", "Producto", "Antes", "Después"], rows)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    audit.record(db, user.upn, "report.monthly", "inventory_snapshot", snap.id, f"Informe {snap.period}")
    db.commit()
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition":
                                      f'attachment; filename="inventario_{snap.period}_v{snap.version}.xlsx"'})


# --------------------------------------------------------------------------- cambios
def _changes_query(change_type, device_id, product_id, since, until, source):
    query = select(InventoryChange)
    if change_type:
        query = query.where(InventoryChange.change_type.in_(change_type.split(",")))
    if device_id:
        query = query.where(InventoryChange.device_id == device_id)
    if product_id:
        query = query.where(InventoryChange.product_id == product_id)
    if since:
        query = query.where(InventoryChange.detected_at >= since)
    if until:
        query = query.where(InventoryChange.detected_at < until)
    if source:
        query = query.where(InventoryChange.source == source)
    return query


def _change(c: InventoryChange) -> dict:
    return {"id": c.id, "type": c.change_type, "device_id": c.device_id,
            "device": c.device.display_name if c.device else None, "product_id": c.product_id,
            "product": c.product.name if c.product else None, "field": c.field, "old": c.old_value,
            "new": c.new_value, "source": c.source, "at": c.detected_at, "run_id": c.run_id}


@router.get("/changes")
def list_changes(page: Page = Depends(), type: str | None = None, device_id: int | None = None,
                 product_id: int | None = None, since: datetime | None = None, until: datetime | None = None,
                 source: str | None = None, db: Session = Depends(get_db), _: CurrentUser = Reader):
    query = _changes_query(type, device_id, product_id, since, until, source)
    rows, total = paginate(db, query, page, {"at": InventoryChange.detected_at}, InventoryChange.id.desc())
    return {"total": total, "items": [_change(r[0]) for r in rows]}


@router.get("/changes/export")
def export_changes(format: str = Query("xlsx", pattern="^(csv|xlsx)$"), type: str | None = None,
                   since: datetime | None = None, until: datetime | None = None, db: Session = Depends(get_db),
                   user: CurrentUser = Reader):
    rows = [_change(c) for c in db.execute(_changes_query(type, None, None, since, until, None)
                                           .order_by(InventoryChange.id.desc()).limit(100_000)).scalars()]
    audit.record(db, user.upn, "export.changes", "inventory_change", None, f"{len(rows)} cambios")
    db.commit()
    return export_response(rows, [("at", "Fecha (UTC)"), ("type", "Tipo"), ("device", "Equipo"),
                                  ("product", "Producto"), ("field", "Campo"), ("old", "Antes"), ("new", "Después"),
                                  ("source", "Fuente"), ("run_id", "Ejecución")], f"cambios_{utcnow():%Y%m%d}", format)


# --------------------------------------------------------------------------- alertas
def _alert(a: Alert) -> dict:
    return {"id": a.id, "type": a.alert_type, "severity": a.severity, "status": a.status, "title": a.title,
            "detail": a.detail, "device_id": a.device_id, "product_id": a.product_id, "run_id": a.run_id,
            "created_at": a.created_at, "last_seen_at": a.last_seen_at, "resolved_at": a.resolved_at,
            "resolved_by": a.resolved_by, "resolution": a.resolution}


@router.get("/alerts")
def list_alerts(page: Page = Depends(), status: str | None = "open,acknowledged", severity: str | None = None,
                type: str | None = None, db: Session = Depends(get_db), _: CurrentUser = Reader):
    query = select(Alert)
    if status:
        query = query.where(Alert.status.in_(status.split(",")))
    if severity:
        query = query.where(Alert.severity.in_(severity.split(",")))
    if type:
        query = query.where(Alert.alert_type.in_(type.split(",")))
    sev_order = func.coalesce(func.nullif(Alert.severity, ""), "")
    rows, total = paginate(db, query, page, {"created_at": Alert.created_at, "severity": sev_order},
                           Alert.id.desc())
    return {"total": total, "items": [_alert(r[0]) for r in rows]}


class AlertAction(BaseModel):
    action: str = Field(..., pattern="^(acknowledge|resolve|reopen)$")
    note: str | None = Field(None, max_length=2000)


@router.post("/alerts/{alert_id}")
def act_alert(alert_id: int, body: AlertAction, db: Session = Depends(get_db), user: CurrentUser = Technician):
    a = db.get(Alert, alert_id)
    if a is None:
        raise not_found("Alerta")
    before = _alert(a)
    if body.action == "acknowledge":
        a.status = "acknowledged"
    elif body.action == "resolve":
        if not (body.note or "").strip():
            raise HTTPException(422, "Indique cómo se ha resuelto")
        a.status, a.resolved_at, a.resolved_by, a.resolution = "resolved", utcnow(), user.upn, body.note
    else:
        a.status, a.resolved_at = "open", None
    audit.record(db, user.upn, f"alert.{body.action}", "alert", a.id, a.title, before={"status": before["status"]},
                 after={"status": a.status, "note": body.note})
    db.commit()
    return _alert(a)


# --------------------------------------------------------------------------- casos de revisión
@router.get("/review-cases")
def list_cases(status: str = "open", db: Session = Depends(get_db), _: CurrentUser = Reader):
    rows = db.execute(select(ReviewCase).where(ReviewCase.status == status).order_by(ReviewCase.id.desc())).scalars()
    return [{"id": c.id, "type": c.case_type, "status": c.status, "summary": c.summary, "data": c.data,
             "run_id": c.run_id, "created_at": c.created_at, "resolved_by": c.resolved_by,
             "resolution": c.resolution} for c in rows]


class CaseResolution(BaseModel):
    action: str = Field(..., pattern="^(dismiss|merge_devices)$")
    note: str = Field(..., min_length=3, max_length=2000)
    keep_device_id: int | None = None
    merge_device_id: int | None = None


@router.post("/review-cases/{case_id}")
def resolve_case(case_id: int, body: CaseResolution, db: Session = Depends(get_db), user: CurrentUser = Technician):
    case = db.get(ReviewCase, case_id)
    if case is None or case.status != "open":
        raise not_found("Caso abierto")
    if body.action == "merge_devices":
        if not (body.keep_device_id and body.merge_device_id) or body.keep_device_id == body.merge_device_id:
            raise HTTPException(422, "Indique el dispositivo que se conserva y el que se fusiona")
        merge_devices(db, body.keep_device_id, body.merge_device_id, user.upn, body.note)
    case.status = "resolved" if body.action == "merge_devices" else "dismissed"
    case.resolved_at, case.resolved_by, case.resolution = utcnow(), user.upn, body.note
    audit.record(db, user.upn, f"review.{body.action}", "review_case", case.id, case.summary,
                 after=body.model_dump())
    db.commit()
    return {"ok": True}


def merge_devices(db: Session, keep_id: int, merge_id: int, actor: str, note: str) -> None:
    """Fusiona `merge` en `keep`: identificadores e instalaciones pasan a `keep`; `merge` queda retirado."""
    keep, merge = db.get(Device, keep_id), db.get(Device, merge_id)
    if keep is None or merge is None:
        raise not_found("Dispositivo")
    db.execute(update(DeviceExternalId).where(DeviceExternalId.device_id == merge.id).values(device_id=keep.id))
    for inst in db.execute(select(SoftwareInstallation).where(SoftwareInstallation.device_id == merge.id)).scalars():
        clash = db.execute(select(SoftwareInstallation).where(
            SoftwareInstallation.device_id == keep.id, SoftwareInstallation.product_id == inst.product_id,
            SoftwareInstallation.version_id == inst.version_id, SoftwareInstallation.source == inst.source,
        )).scalar_one_or_none()
        if clash:
            clash.present = clash.present or inst.present
            inst.present = False
        else:
            inst.device_id = keep.id
    for field in ("serial_number", "manufacturer", "model", "os_name", "os_version", "ram_bytes", "location",
                  "department_id", "primary_user_id"):
        if getattr(keep, field) is None and getattr(merge, field) is not None:
            setattr(keep, field, getattr(merge, field))
    if merge.legacy_equipo_id and not keep.legacy_equipo_id:
        keep.legacy_equipo_id, merge.legacy_equipo_id = merge.legacy_equipo_id, None
    merge.status, merge.retired_at = "retired", utcnow()
    merge.notes = f"{merge.notes or ''}\n[Fusionado en {keep.uid} por {actor}: {note}]".strip()
    db.add(InventoryChange(change_type="device_merged", device_id=keep.id, old_value=str(merge.id),
                           new_value=str(keep.id), source="manual", detected_at=utcnow()))
    db.execute(update(DeviceAssignment).where(DeviceAssignment.device_id == merge.id,
                                              DeviceAssignment.ended_at.is_(None)).values(ended_at=utcnow()))
    audit.record(db, actor, "device.merge", "device", keep.id, f"Fusionado {merge.id} en {keep.id}: {note}")


# --------------------------------------------------------------------------- auditoría
@router.get("/audit")
def list_audit(page: Page = Depends(), entity_type: str | None = None, actor: str | None = None,
               action: str | None = None, since: datetime | None = None, db: Session = Depends(get_db),
               _: CurrentUser = Reader):
    query = select(AuditEvent)
    if entity_type:
        query = query.where(AuditEvent.entity_type == entity_type)
    if actor:
        query = query.where(AuditEvent.actor.ilike(f"%{actor}%"))
    if action:
        query = query.where(AuditEvent.action.startswith(action))
    if since:
        query = query.where(AuditEvent.occurred_at >= since)
    rows, total = paginate(db, query, page, {"at": AuditEvent.occurred_at}, AuditEvent.id.desc())
    return {"total": total, "items": [{"id": e.id, "at": e.occurred_at, "actor": e.actor, "action": e.action,
                                       "entity_type": e.entity_type, "entity_id": e.entity_id, "summary": e.summary,
                                       "before": e.before, "after": e.after, "hash": e.hash} for (e,) in rows]}


@router.get("/audit/verify")
def verify_audit(db: Session = Depends(get_db), _: CurrentUser = Reader):
    return audit.verify_chain(db)


@router.get("/audit/export")
def export_audit(format: str = Query("xlsx", pattern="^(csv|xlsx)$"), since: datetime | None = None,
                 db: Session = Depends(get_db), user: CurrentUser = Reader):
    query = select(AuditEvent).order_by(AuditEvent.id)
    if since:
        query = query.where(AuditEvent.occurred_at >= since)
    rows = [{"id": e.id, "at": e.occurred_at, "actor": e.actor, "action": e.action, "entity": e.entity_type,
             "entity_id": e.entity_id, "summary": e.summary, "hash": e.hash, "prev_hash": e.prev_hash}
            for e in db.execute(query).scalars()]
    audit.record(db, user.upn, "export.audit", "audit_event", None, f"{len(rows)} eventos")
    db.commit()
    return export_response(rows, [("id", "Id"), ("at", "Fecha (UTC)"), ("actor", "Actor"), ("action", "Acción"),
                                  ("entity", "Entidad"), ("entity_id", "Id entidad"), ("summary", "Resumen"),
                                  ("hash", "Hash"), ("prev_hash", "Hash anterior")],
                           f"auditoria_{utcnow():%Y%m%d}", format)


# --------------------------------------------------------------------------- dashboard y salud
@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), _: CurrentUser = Reader):
    s = get_settings()
    now = utcnow()
    stale_limit = now - timedelta(days=s.stale_device_days)
    active = Device.status != "retired"

    def scalar(q):
        return db.execute(q).scalar() or 0

    by_status = dict(db.execute(select(Device.status, func.count()).group_by(Device.status)).all())
    total = scalar(select(func.count()).select_from(Device).where(active))
    current = scalar(select(func.count()).select_from(Device).where(active, Device.last_seen_at >= stale_limit))
    never = scalar(select(func.count()).select_from(Device).where(active, Device.last_seen_at.is_(None)))
    compliance = dict(db.execute(select(Device.compliance_state, func.count()).where(active)
                                 .group_by(Device.compliance_state)).all())
    unassigned = scalar(select(func.count()).select_from(Device).where(active, Device.department_id.is_(None)))
    products = dict(db.execute(select(SoftwareProduct.authorization_status, func.count())
                               .where(SoftwareProduct.merged_into_id.is_(None), SoftwareProduct.is_relevant.is_(True))
                               .group_by(SoftwareProduct.authorization_status)).all())
    installs = scalar(select(func.count()).select_from(SoftwareInstallation).where(SoftwareInstallation.present.is_(True)))
    alerts = dict(db.execute(select(Alert.severity, func.count()).where(Alert.status != "resolved")
                             .group_by(Alert.severity)).all())
    period = now.strftime("%Y-%m")
    snap = db.execute(select(InventorySnapshot).where(InventorySnapshot.period == period)
                      .order_by(InventorySnapshot.version.desc()).limit(1)).scalar_one_or_none()
    recent_runs = db.execute(select(InventoryRun).order_by(InventoryRun.id.desc()).limit(6)).scalars().all()
    recent_changes = db.execute(select(InventoryChange).where(InventoryChange.change_type != "software_new_product")
                                .order_by(InventoryChange.id.desc()).limit(10)).scalars().all()
    top_alerts = db.execute(select(Alert).where(Alert.status != "resolved").order_by(
        (Alert.severity == "critical").desc(), Alert.id.desc()).limit(6)).scalars().all()
    since_30 = now - timedelta(days=30)
    changes_30 = dict(db.execute(select(InventoryChange.change_type, func.count())
                                 .where(InventoryChange.detected_at >= since_30)
                                 .group_by(InventoryChange.change_type)).all())
    dept_rows = db.execute(select(Department.name, func.count(Device.id)).join(
        Device, (Device.department_id == Department.id) & active).group_by(Department.name)
        .order_by(func.count(Device.id).desc())).all()
    hb = db.get(AppSetting, "worker.heartbeat")
    return {
        "generated_at": now,
        "devices": {"total": total, "current": current, "stale": total - current - never, "never_synced": never,
                    "unassigned": unassigned, "by_status": by_status, "compliance": {str(k): v for k, v in compliance.items()}},
        "software": {"products": sum(products.values()), "by_status": products, "installations": installs,
                     "pending": products.get("unclassified", 0) + products.get("pending_review", 0),
                     "unauthorized_products": products.get("unauthorized", 0)},
        "alerts": {"open": sum(alerts.values()), "by_severity": alerts,
                   "top": [_alert(a) for a in top_alerts]},
        "monthly": {"period": period, "snapshot": _snap(snap) if snap else None,
                    "coverage": snapshots.coverage(db, period)},
        "runs": [_run(r) for r in recent_runs],
        "changes": [_change(c) for c in recent_changes],
        "changes_30d": changes_30,
        "departments": [{"name": n, "devices": c} for n, c in dept_rows],
        "worker_heartbeat": hb.value if hb else None,
        "open_review_cases": scalar(select(func.count()).select_from(ReviewCase).where(ReviewCase.status == "open")),
    }


health = APIRouter(tags=["Sistema"])


@health.get("/api/health")
def healthcheck(db: Session = Depends(get_db)):
    """Comprobación de salud sin datos sensibles (para monitorización)."""
    try:
        db.execute(select(1))
        db_ok = True
    except Exception:  # noqa: BLE001
        db_ok = False
    hb = db.get(AppSetting, "worker.heartbeat") if db_ok else None
    worker_ok = None
    if hb and isinstance(hb.value, dict) and hb.value.get("at"):
        worker_ok = utcnow() - datetime.fromisoformat(hb.value["at"]) < timedelta(minutes=10)
    return {"status": "ok" if db_ok else "degraded", "database": db_ok, "worker": worker_ok}
