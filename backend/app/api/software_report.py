"""Excel de software por departamento.

- Una hoja por departamento (y "Sin departamento" si hay equipos sin asignar): cada programa
  instalado en sus equipos con el número de equipos y cuáles son. Las filas de programas que
  solo están en un equipo del departamento se marcan en amarillo.
- Última hoja: inventario de equipos (las mismas columnas que la exportación de dispositivos).
"""
from __future__ import annotations

import io
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import false, select
from sqlalchemy.orm import Session

from app.api.common import _cell, _csv_safe
from app.models import Department, Device, SoftwareInstallation, SoftwareProduct, SoftwareVersion

CATEGORY_LABEL = {"application": "Aplicación", "system_component": "Componente del sistema", "driver": "Controlador",
                  "runtime": "Runtime / librería", "update": "Actualización", "unknown": "Desconocido"}
AUTH_LABEL = {"authorized": "Autorizado", "unauthorized": "No autorizado", "pending_review": "Pendiente de revisión",
              "exception": "Excepción aprobada", "unclassified": "No clasificado"}
DEVICE_STATUS_LABEL = {"active": "En uso", "in_stock": "En almacén", "repair": "En reparación",
                       "retired": "Dado de baja", "lost": "Perdido / robado"}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
SINGLE_FILL = PatternFill("solid", fgColor="FFF2CC")  # amarillo claro: programa en un solo equipo
SOFTWARE_COLUMNS = [("Programa", 45), ("Editor", 30), ("Categoría", 22), ("Autorización", 22), ("Versiones", 28),
                    ("Nº equipos", 11), ("Equipos", 70)]
NO_DEPARTMENT = "Sin departamento"


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name).strip("'")[:31] or "Hoja"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def _header(ws, row: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, 1):
        cell = ws.cell(row, col, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def build_workbook(db: Session, product_ids: set[int], department_id: int | None,
                   device_rows: list[dict], device_columns: list[tuple[str, str]], generated: str) -> bytes:
    rows = db.execute(
        select(Device.id, Device.department_id, Device.hostname, Device.serial_number,
               SoftwareInstallation.product_id, SoftwareVersion.version)
        .join(SoftwareInstallation, SoftwareInstallation.device_id == Device.id)
        .join(SoftwareVersion, SoftwareVersion.id == SoftwareInstallation.version_id)
        .where(SoftwareInstallation.present.is_(True), Device.status != "retired",
               SoftwareInstallation.product_id.in_(product_ids) if product_ids else false())
    ).all()
    # departamento -> producto -> {equipo: nombre}, y versiones
    per_dept: dict[int | None, dict[int, dict[int, str]]] = defaultdict(lambda: defaultdict(dict))
    versions: dict[tuple[int | None, int], set[str]] = defaultdict(set)
    for dev_id, dept_id, hostname, serial, product_id, version in rows:
        if department_id and dept_id != department_id:
            continue
        per_dept[dept_id][product_id][dev_id] = hostname or serial or f"#{dev_id}"
        versions[(dept_id, product_id)].add(version or "?")

    products = {p.id: p for p in db.execute(select(SoftwareProduct).where(
        SoftwareProduct.id.in_({pid for prods in per_dept.values() for pid in prods}))).scalars()}
    departments = {d.id: d.name for d in db.execute(select(Department)).scalars()}
    dept_devices: dict[int | None, int] = defaultdict(int)
    for d in db.execute(select(Device.department_id).where(Device.status != "retired")).scalars():
        dept_devices[d] += 1

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    order = sorted(per_dept, key=lambda d: (d is None, (departments.get(d) or "").casefold()))
    if department_id and department_id not in per_dept:
        order = [department_id]
    for dept_id in order:
        name = departments.get(dept_id, NO_DEPARTMENT) if dept_id is not None else NO_DEPARTMENT
        ws = wb.create_sheet(_sheet_title(name, used))
        prods = per_dept.get(dept_id, {})
        single = sum(len(devs) == 1 for devs in prods.values())
        ws.cell(1, 1, name).font = Font(bold=True, size=14)
        ws.cell(2, 1, f"{dept_devices.get(dept_id, 0)} equipos · {len(prods)} programas · "
                      f"{single} en un solo equipo (en amarillo) · generado {generated}").font = Font(italic=True,
                                                                                                     color="6B7280")
        _header(ws, 4, [t for t, _ in SOFTWARE_COLUMNS])
        ordered = sorted(prods.items(), key=lambda kv: (-len(kv[1]), products[kv[0]].name.casefold()))
        for r, (pid, devs) in enumerate(ordered, 5):
            p = products[pid]
            values = [p.name, p.publisher, CATEGORY_LABEL.get(p.category, p.category),
                      AUTH_LABEL.get(p.authorization_status, p.authorization_status),
                      ", ".join(sorted(versions[(dept_id, pid)])), len(devs),
                      ", ".join(sorted(devs.values(), key=str.casefold))]
            for c, value in enumerate(values, 1):
                cell = ws.cell(r, c, _csv_safe(value))  # openpyxl trataría "=..." como fórmula
                if len(devs) == 1:
                    cell.fill = SINGLE_FILL
        for c, (_, width) in enumerate(SOFTWARE_COLUMNS, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        last = 4 + len(ordered)
        ws.auto_filter.ref = f"A4:{get_column_letter(len(SOFTWARE_COLUMNS))}{max(last, 5)}"
        ws.freeze_panes = "B5"

    ws = wb.create_sheet(_sheet_title("Inventario de equipos", used))
    _header(ws, 1, [t for _, t in device_columns])
    for r, row in enumerate(device_rows, 2):
        row = {**row, "status": DEVICE_STATUS_LABEL.get(row.get("status"), row.get("status"))}
        for c, (key, _) in enumerate(device_columns, 1):
            ws.cell(r, c, _csv_safe(_cell(row.get(key))))
    for c, (_, title) in enumerate(device_columns, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(40, len(title) + 4))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(device_columns))}{max(len(device_rows) + 1, 2)}"
    ws.freeze_panes = "B2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
