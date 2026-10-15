from __future__ import annotations

import time

import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.config import Settings, get_settings
from app.core.security import CurrentUser, get_current_user
from app.main import create_app
from app.models import Alert, AuditEvent, Device, InventoryRun
from app.services import runs
from app.services.audit import verify_chain

PANDA = ("Nombre\tEditor\tFecha de instalación\tTamaño\tVersión\n"
         "AnyDesk\tAnyDesk Software GmbH\t10/03/2026\t2 MB\t9.0.10\n"
         "Google Chrome\tGoogle LLC\t01/02/2026\t100 MB\t129.0.1\n")


@pytest.fixture()
def app(engine):
    return create_app()


def as_role(app, role: str):
    app.dependency_overrides[get_current_user] = lambda: CurrentUser(upn=f"{role}@asserta.test", name=role, role=role)


@pytest.fixture()
def client(app):
    as_role(app, "admin")
    return TestClient(app)


def drain(db):
    while (run := runs.claim_next(db)) is not None:
        runs.execute_run(db, run)


def test_health_and_public_config(client):
    assert client.get("/api/health").json()["database"] is True
    cfg = client.get("/api/config/public").json()
    assert cfg["auth_mode"] == "dev" and "secret" not in str(cfg).lower()


def test_rbac_enforced_in_backend(app, client):
    as_role(app, "auditor")
    assert client.get("/api/devices").status_code == 200
    assert client.post("/api/departments", json={"name": "IT"}).status_code == 403
    assert client.post("/api/devices", json={"hostname": "X"}).status_code == 403
    as_role(app, "technician")
    assert client.post("/api/departments", json={"name": "IT"}).status_code == 201
    assert client.put("/api/settings/schedule", json={"enabled": True, "day": 1, "hour": 6,
                                                      "sources": ["intune"]}).status_code == 403
    as_role(app, "admin")
    assert client.put("/api/settings/schedule", json={"enabled": True, "day": 2, "hour": 5,
                                                      "sources": ["intune"]}).status_code == 200


def _entra_settings() -> Settings:
    return Settings(environment="test", auth_mode="entra", entra_tenant_id="tid", entra_api_client_id="api-id")


def _token(key, **claims):
    now = int(time.time())
    body = {"iss": "https://login.microsoftonline.com/tid/v2.0", "aud": "api-id", "iat": now, "nbf": now,
            "exp": now + 600, "tid": "tid", "oid": "o1", "preferred_username": "Ana@Asserta.test",
            "name": "Ana", "roles": ["Inventario.Tecnico"]}
    body.update(claims)
    return jwt.encode(body, key, algorithm="RS256")


def test_entra_token_validation(app):
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    other = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    app.dependency_overrides[get_settings] = _entra_settings
    app.state.jwt_key_resolver = lambda token: key.public_key()
    c = TestClient(app)
    assert c.get("/api/me").status_code == 401
    ok = c.get("/api/me", headers={"Authorization": f"Bearer {_token(key)}"})
    assert ok.status_code == 200 and ok.json()["role"] == "technician" and ok.json()["upn"] == "ana@asserta.test"
    assert c.get("/api/me", headers={"Authorization": f"Bearer {_token(other)}"}).status_code == 401  # firma
    assert c.get("/api/me", headers={"Authorization": f"Bearer {_token(key, aud='otra')}"}).status_code == 401
    assert c.get("/api/me", headers={"Authorization": f"Bearer {_token(key, exp=1)}"}).status_code == 401
    assert c.get("/api/me", headers={"Authorization": f"Bearer {_token(key, roles=[])}"}).status_code == 403
    assert c.post("/api/departments", json={"name": "X"},
                  headers={"Authorization": f"Bearer {_token(key, roles=['Inventario.Auditor'])}"}).status_code == 403


def test_production_requires_entra():
    with pytest.raises(ValueError):
        Settings(environment="production", auth_mode="dev", database_url="mysql+pymysql://u:p@h/db")


def test_department_lifecycle_and_bulk_move(client, db):
    it = client.post("/api/departments", json={"name": "IT", "description": "Sistemas"}).json()
    data = client.post("/api/departments", json={"name": "Data Science/Analytics"}).json()
    assert data["code"] == "data_science_analytics"
    assert client.post("/api/departments", json={"name": "it"}).status_code == 409
    d1 = client.post("/api/devices", json={"hostname": "PC-1", "department_id": it["id"]}).json()
    d2 = client.post("/api/devices", json={"hostname": "PC-2", "department_id": it["id"]}).json()
    # No se puede desactivar con dispositivos asignados
    assert client.patch(f"/api/departments/{it['id']}", json={"active": False}).status_code == 409
    r = client.post("/api/devices/bulk-assign", json={"device_ids": [d1["id"], d2["id"]],
                                                      "department_id": data["id"], "reason": "Reorganización"})
    assert r.json() == {"moved": 2, "unchanged": 0}
    hist = client.get(f"/api/devices/{d1['id']}/assignments").json()
    assert [h["department"] for h in hist] == ["Data Science/Analytics", "IT"]
    assert hist[1]["ended_at"] is not None  # historial preservado
    assert client.patch(f"/api/departments/{it['id']}", json={"active": False}).status_code == 200
    depts = client.get("/api/departments").json()["items"]
    assert {d["name"]: d["devices"] for d in depts} == {"Data Science/Analytics": 2, "IT": 0}
    assert verify_chain(db)["valid"]


def test_device_crud_retire_and_filters(client):
    d = client.post("/api/devices", json={"hostname": "SRV-01", "serial_number": "ABC12345",
                                          "device_type_id": 3, "primary_user_upn": "Ana@Asserta.test"}).json()
    assert d["primary_user_upn"] == "ana@asserta.test" and d["device_type"] == "Servidor"
    assert d["freshness"] == "never_synced"
    upd = client.patch(f"/api/devices/{d['id']}", json={"location": "CPD", "model": "PowerEdge"}).json()
    assert upd["field_sources"]["model"]["source"] == "manual"
    assert client.get("/api/devices", params={"q": "ABC123"}).json()["total"] == 1
    assert client.post(f"/api/devices/{d['id']}/retire", json={"reason": "Fin de vida"}).status_code == 200
    assert client.get("/api/devices").json()["total"] == 0  # retirados fuera por defecto
    assert client.get("/api/devices", params={"freshness": "retired"}).json()["total"] == 1
    assert client.patch(f"/api/devices/{d['id']}", json={"location": "X"}).status_code == 409
    hist = client.get(f"/api/devices/{d['id']}/history").json()
    assert {"device_added", "device_retired"} <= {c["type"] for c in hist["changes"]}
    assert client.post("/api/devices", json={}).status_code == 422


def test_panda_upload_flow_and_software_queries(client, db):
    dev = client.post("/api/devices", json={"hostname": "PC-7"}).json()
    r = client.post("/api/runs/upload", data={"source": "panda", "device_id": str(dev["id"])},
                    files={"file": ("export.txt", PANDA.encode("utf-8"), "text/plain")})
    assert r.status_code == 202
    drain(db)
    run = client.get(f"/api/runs/{r.json()['id']}").json()
    assert run["status"] == "succeeded", run["error"]
    sw = client.get(f"/api/devices/{dev['id']}/software").json()
    assert {i["name"] for i in sw["items"]} == {"AnyDesk", "Google Chrome"}
    assert all(i["authorization_status"] == "unclassified" for i in sw["items"])  # nunca autoriza solo
    # ¿Qué dispositivos tienen AnyDesk?
    anydesk = next(i for i in sw["items"] if i["name"] == "AnyDesk")
    inst = client.get(f"/api/software/{anydesk['product_id']}/installations").json()
    assert [i["device"] for i in inst["items"]] == ["PC-7"]
    # Prohibirlo globalmente genera alerta; una excepción para el equipo la resuelve
    assert client.post("/api/policies", json={"product_id": anydesk["product_id"], "scope": "global",
                                              "status": "unauthorized"}).status_code == 201
    alerts = client.get("/api/alerts", params={"type": "unauthorized_software"}).json()
    assert alerts["total"] == 1
    bad = client.post("/api/policies", json={"product_id": anydesk["product_id"], "scope": "device",
                                             "device_id": dev["id"], "status": "exception"})
    assert bad.status_code == 422  # sin justificación
    ok = client.post("/api/policies", json={"product_id": anydesk["product_id"], "scope": "device",
                                            "device_id": dev["id"], "status": "exception",
                                            "justification": "Soporte remoto a proveedor",
                                            "approved_by": "ciso@asserta.test", "valid_until": "2099-01-01T00:00:00"})
    assert ok.status_code == 201
    assert client.get("/api/alerts", params={"type": "unauthorized_software"}).json()["total"] == 0
    listed = client.get("/api/software", params={"status": "unauthorized"}).json()
    assert listed["total"] == 1 and listed["items"][0]["devices"] == 1
    assert client.get("/api/software", params={"department_id": 999}).json()["total"] == 0


def test_snapshot_generate_review_compare_and_report(client, db):
    dev = client.post("/api/devices", json={"hostname": "PC-9"}).json()
    client.post("/api/runs/upload", data={"source": "panda", "device_id": str(dev["id"])},
                files={"file": ("x.txt", PANDA.encode(), "text/plain")})
    drain(db)
    s1 = client.post("/api/snapshots", json={"period": "2026-09"}).json()
    assert s1["totals"]["installations"] == 2
    # Intune nunca sincronizó en el periodo: el inventario no está completo
    assert s1["is_complete"] is False
    r = client.post(f"/api/snapshots/{s1['id']}/review", json={"approve": True})
    assert r.status_code == 409
    r = client.post(f"/api/snapshots/{s1['id']}/review", json={"approve": True, "accept_incomplete": True,
                                                               "notes": "Intune pendiente de credenciales"})
    assert r.json()["status"] == "approved"
    upd = PANDA.replace("129.0.1", "130.0.0").replace("AnyDesk\tAnyDesk Software GmbH\t10/03/2026\t2 MB\t9.0.10\n", "")
    client.post("/api/runs/upload", data={"source": "panda", "device_id": str(dev["id"])},
                files={"file": ("x.txt", upd.encode(), "text/plain")})
    drain(db)
    s2 = client.post("/api/snapshots", json={"period": "2026-10"}).json()
    diff = client.get("/api/snapshots/compare", params={"base": s1["id"], "target": s2["id"]}).json()
    assert [x["product"] for x in diff["installs_removed"]] == ["AnyDesk"]
    assert diff["versions_changed"][0]["from"] == ["129.0.1"] and diff["versions_changed"][0]["to"] == ["130.0.0"]
    rep = client.get(f"/api/snapshots/{s2['id']}/report")
    assert rep.status_code == 200
    import io
    wb = load_workbook(io.BytesIO(rep.content))
    assert "Cambios vs 2026-09" in wb.sheetnames and "Cobertura" in wb.sheetnames
    changes = client.get("/api/changes", params={"type": "software_version_changed"}).json()
    assert changes["total"] == 1


def test_exports_neutralize_formulas(client):
    client.post("/api/devices", json={"hostname": "=cmd|'/c calc'!A1"})
    csv_text = client.get("/api/devices/export", params={"format": "csv"}).content.decode("utf-8-sig")
    assert "'=cmd" in csv_text
    import io
    wb = load_workbook(io.BytesIO(client.get("/api/devices/export", params={"format": "xlsx"}).content))
    assert wb.active["A2"].value.startswith("'=")


def test_manual_sync_requires_graph_credentials(client):
    r = client.post("/api/runs", json={"source": "intune"})
    assert r.status_code == 409 and "credenciales" in r.json()["detail"]
    assert client.post("/api/runs", json={"source": "panda"}).status_code == 422


def test_failed_run_retry_and_dashboard(client, db):
    r = client.post("/api/runs/upload", data={"source": "panda"},
                    files={"file": ("SIN-CONTENIDO.txt", b"basura sin formato", "text/plain")})
    drain(db)
    run = client.get(f"/api/runs/{r.json()['id']}").json()
    assert run["status"] == "failed" and "[data]" in run["error"]
    assert client.post(f"/api/runs/{run['id']}/retry").status_code == 202
    dash = client.get("/api/dashboard").json()
    assert dash["alerts"]["open"] >= 1 and dash["monthly"]["period"]
    assert db.execute(select(Alert).where(Alert.alert_type == "sync_failed")).scalars().first()


def test_audit_endpoints(client, db):
    client.post("/api/departments", json={"name": "Gerencia"})
    events = client.get("/api/audit").json()
    assert events["items"][0]["action"] == "department.create"
    assert client.get("/api/audit/verify").json()["valid"] is True
    # Manipulación directa en BD -> la verificación lo detecta
    ev = db.execute(select(AuditEvent)).scalars().first()
    ev.summary = "manipulado"
    db.commit()
    assert client.get("/api/audit/verify").json()["valid"] is False


def test_review_case_merge_devices(client, db):
    a = client.post("/api/devices", json={"hostname": "PC-A", "serial_number": "SER-0001"}).json()
    b = client.post("/api/devices", json={"hostname": "PC-A-OLD", "location": "Planta 2"}).json()
    from app.models import ReviewCase

    db.add(ReviewCase(case_type="device_match", dedup_key="t", summary="duplicado", data={}))
    db.commit()
    case = client.get("/api/review-cases").json()[0]
    r = client.post(f"/api/review-cases/{case['id']}", json={"action": "merge_devices", "note": "mismo equipo",
                                                            "keep_device_id": a["id"], "merge_device_id": b["id"]})
    assert r.status_code == 200
    assert db.get(Device, b["id"]).status == "retired"
    assert db.get(Device, a["id"]).location == "Planta 2"


def test_validation_errors_are_readable(client):
    r = client.post("/api/departments", json={"name": "A"})
    assert r.status_code == 422 and isinstance(r.json()["detail"], str)


def test_queue_status_visible(client, db):
    client.post("/api/runs/upload", data={"source": "panda"},
                files={"file": ("PC-55.txt", PANDA.encode(), "text/plain")})
    assert client.get("/api/runs", params={"status": "queued"}).json()["total"] == 1
    drain(db)
    assert db.execute(select(InventoryRun.status)).scalar_one() == "succeeded"
    # Panda sin dispositivo previo: crea la ficha por hostname
    assert client.get("/api/devices", params={"q": "PC-55"}).json()["total"] == 1


def test_create_device_with_explicit_status(client):
    r = client.post("/api/devices", json={"hostname": "PC-STOCK", "status": "in_stock"})
    assert r.status_code == 201 and r.json()["status"] == "in_stock"


def test_local_mode_only_on_loopback():
    Settings(environment="production", auth_mode="local", bind_host="127.0.0.1")  # válido, SQLite incluido
    with pytest.raises(ValueError):
        Settings(environment="production", auth_mode="local", bind_host="0.0.0.0")
    with pytest.raises(ValueError):
        Settings(environment="production", auth_mode="entra", entra_tenant_id="t", entra_api_client_id="a",
                 database_url="sqlite:///x.db")


def test_local_mode_rejects_remote_clients():
    from types import SimpleNamespace

    from fastapi import HTTPException

    s = Settings(environment="production", auth_mode="local")
    local = SimpleNamespace(client=SimpleNamespace(host="127.0.0.1"), headers={})
    assert get_current_user(local, db=None, settings=s).role == "admin"
    remote = SimpleNamespace(client=SimpleNamespace(host="192.168.1.50"), headers={})
    with pytest.raises(HTTPException) as err:
        get_current_user(remote, db=None, settings=s)
    assert err.value.status_code == 403
