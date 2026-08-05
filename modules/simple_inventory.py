from __future__ import annotations

import re
from hashlib import sha256
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text

from utils.panda_parser import parse_panda_text
from utils.normalizer import normalize_nombre


PERIODO_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")


def current_period(today: date | None = None) -> str:
    value = today or date.today()
    return f"{value.year:04d}-{value.month:02d}"


def validate_periodo(periodo: str) -> str:
    value = str(periodo or "").strip()
    if not PERIODO_RE.fullmatch(value):
        raise ValueError("El periodo debe tener formato YYYY-MM.")
    return value


def _engine():
    from database.connection import get_engine

    return get_engine()


def confirm_software_import(equipo_id: int, periodo: str, rows: list[dict], raw_text: str, db=None) -> int:
    if db is None:
        with _engine().begin() as conn:
            return confirm_software_import(equipo_id, periodo, rows, raw_text, db=conn)
    periodo = validate_periodo(periodo)
    if not rows:
        raise ValueError("No hay programas validos para guardar.")

    equipo = db.execute(
        text(
            """
            SELECT id, departamento_id
            FROM equipos
            WHERE id = :equipo_id
              AND activo = TRUE
            """
        ),
        {"equipo_id": equipo_id},
    ).mappings().first()
    if not equipo:
        raise ValueError("Equipo activo no encontrado.")

    parsed = parse_panda_text(raw_text)
    departamento_id = int(equipo["departamento_id"])
    db.execute(
        text(
            """
            UPDATE software_importaciones
            SET estado = 'superseded'
            WHERE equipo_id = :equipo_id
              AND periodo = :periodo
              AND estado = 'confirmed'
            """
        ),
        {"equipo_id": equipo_id, "periodo": periodo},
    )
    result = db.execute(
        text(
            """
            INSERT INTO software_importaciones (
                equipo_id, departamento_id, periodo, fecha_importacion,
                raw_hash, raw_text, estado, n_programas, n_errores, origen
            )
            VALUES (
                :equipo_id, :departamento_id, :periodo, NOW(),
                :raw_hash, :raw_text, 'confirmed', :n_programas, 0, :origen
            )
            """
        ),
        {
            "equipo_id": equipo_id,
            "departamento_id": departamento_id,
            "periodo": periodo,
            "raw_hash": parsed.raw_hash,
            "raw_text": raw_text,
            "n_programas": len(rows),
            "origen": parsed.mode,
        },
    )
    importacion_id = int(result.lastrowid)
    for row in rows:
        db.execute(
            text(
                """
                INSERT INTO software_instalado (
                    importacion_id, equipo_id, departamento_id, periodo, source_row,
                    nombre, nombre_norm, fabricante, fabricante_norm,
                    fecha_instalacion, tamano, version, raw_line
                )
                VALUES (
                    :importacion_id, :equipo_id, :departamento_id, :periodo, :source_row,
                    :nombre, :nombre_norm, :fabricante, :fabricante_norm,
                    :fecha_instalacion, :tamano, :version, :raw_line
                )
                """
            ),
            {
                "importacion_id": importacion_id,
                "equipo_id": equipo_id,
                "departamento_id": departamento_id,
                "periodo": periodo,
                "source_row": row.get("row_number"),
                "nombre": row.get("nombre"),
                "nombre_norm": row.get("nombre_norm"),
                "fabricante": row.get("fabricante"),
                "fabricante_norm": row.get("fabricante_norm"),
                "fecha_instalacion": row.get("fecha_instalacion"),
                "tamano": row.get("tamano"),
                "version": row.get("version"),
                "raw_line": row.get("raw_line"),
            },
        )
    return importacion_id


def software_por_departamento(periodo: str, departamento_id: int, db=None) -> list[dict]:
    if db is None:
        with _engine().connect() as conn:
            return software_por_departamento(periodo, departamento_id, db=conn)
    periodo = validate_periodo(periodo)
    rows = db.execute(
        text(
            """
            SELECT
                MIN(si.nombre) AS nombre,
                si.nombre_norm,
                GROUP_CONCAT(DISTINCT NULLIF(si.fabricante, '') ORDER BY si.fabricante SEPARATOR ', ') AS fabricantes,
                GROUP_CONCAT(DISTINCT NULLIF(si.version, '') ORDER BY si.version SEPARATOR ', ') AS versiones,
                COUNT(DISTINCT si.equipo_id) AS n_equipos,
                GROUP_CONCAT(DISTINCT e.nombre ORDER BY e.nombre SEPARATOR ', ') AS equipos
            FROM software_instalado si
            JOIN software_importaciones imp ON imp.id = si.importacion_id
            JOIN equipos e ON e.id = si.equipo_id
            WHERE imp.estado = 'confirmed'
              AND imp.periodo = :periodo
              AND imp.departamento_id = :departamento_id
            GROUP BY si.nombre_norm
            ORDER BY n_equipos DESC, nombre
            """
        ),
        {"periodo": periodo, "departamento_id": departamento_id},
    ).mappings().all()
    return [dict(row) for row in rows]


def equipos_estado_mensual(periodo: str, departamento_id: int, db=None, solo_activos: bool = True) -> list[dict]:
    if db is None:
        with _engine().connect() as conn:
            return equipos_estado_mensual(periodo, departamento_id, db=conn, solo_activos=solo_activos)
    periodo = validate_periodo(periodo)
    active_filter = "AND e.activo = TRUE" if solo_activos else ""
    rows = db.execute(
        text(
            f"""
            SELECT
                e.id,
                e.nombre,
                e.notas AS usuario,
                e.activo,
                d.nombre AS departamento,
                imp.id AS importacion_id,
                imp.fecha_importacion,
                imp.n_programas
            FROM equipos e
            JOIN departamentos d ON d.id = e.departamento_id
            LEFT JOIN software_importaciones imp
              ON imp.equipo_id = e.id
             AND imp.periodo = :periodo
             AND imp.estado = 'confirmed'
            WHERE e.departamento_id = :departamento_id
              {active_filter}
            ORDER BY imp.id IS NULL DESC, e.nombre
            """
        ),
        {"periodo": periodo, "departamento_id": departamento_id},
    ).mappings().all()
    result = []
    for row in rows:
        item = dict(row)
        if not item.get("activo"):
            item["estado_importacion"] = "Inactivo"
        else:
            item["estado_importacion"] = "Importado" if item.get("importacion_id") else "Pendiente"
        item["fecha_importacion_str"] = str(item.get("fecha_importacion") or "")[:19]
        item["n_programas"] = item.get("n_programas") or 0
        item["usuario"] = item.get("usuario") or ""
        result.append(item)
    return result


def software_dispositivos_mensual(periodo: str, departamento_id: int, nombre_norm: str, db=None) -> list[dict]:
    if db is None:
        with _engine().connect() as conn:
            return software_dispositivos_mensual(periodo, departamento_id, nombre_norm, db=conn)
    periodo = validate_periodo(periodo)
    rows = db.execute(
        text(
            """
            SELECT
                e.id,
                e.nombre,
                e.notas AS usuario,
                e.activo,
                EXISTS (
                    SELECT 1
                    FROM software_instalado si
                    JOIN software_importaciones imp ON imp.id = si.importacion_id
                    WHERE imp.estado = 'confirmed'
                      AND imp.periodo = :periodo
                      AND si.equipo_id = e.id
                      AND si.departamento_id = :departamento_id
                      AND si.nombre_norm = :nombre_norm
                ) AS instalado
            FROM equipos e
            WHERE e.departamento_id = :departamento_id
              AND e.activo = TRUE
            ORDER BY e.nombre
            """
        ),
        {"periodo": periodo, "departamento_id": departamento_id, "nombre_norm": nombre_norm},
    ).mappings().all()
    return [{**dict(row), "instalado": bool(row["instalado"])} for row in rows]


def _program_payload(data: dict) -> dict:
    nombre = str(data.get("nombre") or "").strip()
    if not nombre:
        raise ValueError("El nombre del programa es obligatorio.")
    fabricante = str(data.get("fabricante") or "").strip()
    version = str(data.get("version") or "").strip()
    tamano = str(data.get("tamano") or "").strip()
    return {
        "nombre": nombre,
        "nombre_norm": normalize_nombre(nombre),
        "fabricante": fabricante or None,
        "fabricante_norm": normalize_nombre(fabricante) if fabricante else None,
        "version": version or None,
        "tamano": tamano or None,
    }


def _confirmed_import_ids(db, periodo: str, departamento_id: int, nombre_norm: str) -> list[int]:
    rows = db.execute(
        text(
            """
            SELECT DISTINCT imp.id
            FROM software_importaciones imp
            JOIN software_instalado si ON si.importacion_id = imp.id
            WHERE imp.periodo = :periodo
              AND imp.departamento_id = :departamento_id
              AND imp.estado = 'confirmed'
              AND si.nombre_norm = :nombre_norm
            """
        ),
        {"periodo": periodo, "departamento_id": departamento_id, "nombre_norm": nombre_norm},
    ).mappings().all()
    return [int(row["id"]) for row in rows]


def _refresh_program_count(db, importacion_id: int) -> None:
    db.execute(
        text(
            """
            UPDATE software_importaciones
            SET n_programas = (
                SELECT COUNT(*)
                FROM software_instalado
                WHERE importacion_id = :importacion_id
            )
            WHERE id = :importacion_id
            """
        ),
        {"importacion_id": importacion_id},
    )


def _ensure_manual_import(db, equipo_id: int, departamento_id: int, periodo: str) -> int:
    row = db.execute(
        text(
            """
            SELECT id
            FROM software_importaciones
            WHERE equipo_id = :equipo_id
              AND periodo = :periodo
              AND estado = 'confirmed'
            ORDER BY id DESC
            LIMIT 1
            """
        ),
        {"equipo_id": equipo_id, "periodo": periodo},
    ).mappings().first()
    if row:
        return int(row["id"])

    digest = sha256(f"manual:{equipo_id}:{periodo}".encode("utf-8")).hexdigest()
    result = db.execute(
        text(
            """
            INSERT INTO software_importaciones (
                equipo_id, departamento_id, periodo, fecha_importacion,
                raw_hash, raw_text, estado, n_programas, n_errores, origen, notas
            )
            VALUES (
                :equipo_id, :departamento_id, :periodo, NOW(),
                :raw_hash, '', 'confirmed', 0, 0, 'manual', 'Ajuste manual desde Inventario'
            )
            """
        ),
        {
            "equipo_id": equipo_id,
            "departamento_id": departamento_id,
            "periodo": periodo,
            "raw_hash": digest,
        },
    )
    return int(result.lastrowid)


def _next_source_row(db, importacion_id: int) -> int:
    value = db.execute(
        text(
            """
            SELECT COALESCE(MAX(source_row), 0) + 1
            FROM software_instalado
            WHERE importacion_id = :importacion_id
            """
        ),
        {"importacion_id": importacion_id},
    ).scalar()
    return int(value or 1)


def actualizar_software_mensual(
    periodo: str,
    departamento_id: int,
    nombre_norm: str,
    data: dict,
    db=None,
) -> None:
    if db is None:
        with _engine().begin() as conn:
            return actualizar_software_mensual(periodo, departamento_id, nombre_norm, data, db=conn)
    periodo = validate_periodo(periodo)
    payload = _program_payload(data)
    import_ids = _confirmed_import_ids(db, periodo, departamento_id, nombre_norm)
    db.execute(
        text(
            """
            UPDATE software_instalado si
            JOIN software_importaciones imp ON imp.id = si.importacion_id
            SET
                si.nombre = :nombre,
                si.nombre_norm = :new_nombre_norm,
                si.fabricante = :fabricante,
                si.fabricante_norm = :fabricante_norm,
                si.version = :version,
                si.tamano = COALESCE(:tamano, si.tamano)
            WHERE imp.periodo = :periodo
              AND imp.departamento_id = :departamento_id
              AND imp.estado = 'confirmed'
              AND si.nombre_norm = :old_nombre_norm
            """
        ),
        {
            **payload,
            "new_nombre_norm": payload["nombre_norm"],
            "old_nombre_norm": nombre_norm,
            "periodo": periodo,
            "departamento_id": departamento_id,
        },
    )
    for importacion_id in import_ids:
        _refresh_program_count(db, importacion_id)


def eliminar_software_mensual(periodo: str, departamento_id: int, nombre_norm: str, db=None) -> None:
    if db is None:
        with _engine().begin() as conn:
            return eliminar_software_mensual(periodo, departamento_id, nombre_norm, db=conn)
    periodo = validate_periodo(periodo)
    import_ids = _confirmed_import_ids(db, periodo, departamento_id, nombre_norm)
    db.execute(
        text(
            """
            DELETE si
            FROM software_instalado si
            JOIN software_importaciones imp ON imp.id = si.importacion_id
            WHERE imp.periodo = :periodo
              AND imp.departamento_id = :departamento_id
              AND imp.estado = 'confirmed'
              AND si.nombre_norm = :nombre_norm
            """
        ),
        {"periodo": periodo, "departamento_id": departamento_id, "nombre_norm": nombre_norm},
    )
    for importacion_id in import_ids:
        _refresh_program_count(db, importacion_id)


def set_software_dispositivos_mensual(
    periodo: str,
    departamento_id: int,
    old_nombre_norm: str | None,
    data: dict,
    equipo_ids: list[int],
    db=None,
) -> None:
    if db is None:
        with _engine().begin() as conn:
            return set_software_dispositivos_mensual(
                periodo, departamento_id, old_nombre_norm, data, equipo_ids, db=conn
            )
    periodo = validate_periodo(periodo)
    payload = _program_payload(data)
    selected_ids = {int(equipo_id) for equipo_id in equipo_ids}
    target_norm = old_nombre_norm or payload["nombre_norm"]
    touched_imports = set(_confirmed_import_ids(db, periodo, departamento_id, target_norm))

    if old_nombre_norm and old_nombre_norm != payload["nombre_norm"]:
        actualizar_software_mensual(periodo, departamento_id, old_nombre_norm, payload, db=db)
        target_norm = payload["nombre_norm"]

    rows = db.execute(
        text(
            """
            SELECT id
            FROM equipos
            WHERE departamento_id = :departamento_id
              AND activo = TRUE
            """
        ),
        {"departamento_id": departamento_id},
    ).mappings().all()
    active_ids = {int(row["id"]) for row in rows}
    selected_ids &= active_ids

    for equipo_id in sorted(active_ids - selected_ids):
        import_rows = db.execute(
            text(
                """
                SELECT DISTINCT imp.id
                FROM software_importaciones imp
                JOIN software_instalado si ON si.importacion_id = imp.id
                WHERE imp.equipo_id = :equipo_id
                  AND imp.periodo = :periodo
                  AND imp.estado = 'confirmed'
                  AND si.nombre_norm = :nombre_norm
                """
            ),
            {"equipo_id": equipo_id, "periodo": periodo, "nombre_norm": target_norm},
        ).mappings().all()
        for row in import_rows:
            touched_imports.add(int(row["id"]))
        db.execute(
            text(
                """
                DELETE si
                FROM software_instalado si
                JOIN software_importaciones imp ON imp.id = si.importacion_id
                WHERE imp.equipo_id = :equipo_id
                  AND imp.periodo = :periodo
                  AND imp.estado = 'confirmed'
                  AND si.nombre_norm = :nombre_norm
                """
            ),
            {"equipo_id": equipo_id, "periodo": periodo, "nombre_norm": target_norm},
        )

    raw_line = "\t".join(
        [
            payload["nombre"],
            payload.get("fabricante") or "",
            "",
            payload.get("tamano") or "",
            payload.get("version") or "",
        ]
    )
    for equipo_id in sorted(selected_ids):
        importacion_id = _ensure_manual_import(db, equipo_id, departamento_id, periodo)
        touched_imports.add(importacion_id)
        existing = db.execute(
            text(
                """
                SELECT si.id
                FROM software_instalado si
                WHERE si.importacion_id = :importacion_id
                  AND si.nombre_norm = :nombre_norm
                LIMIT 1
                """
            ),
            {"importacion_id": importacion_id, "nombre_norm": target_norm},
        ).mappings().first()
        if existing:
            db.execute(
                text(
                    """
                    UPDATE software_instalado
                    SET nombre = :nombre,
                        nombre_norm = :nombre_norm,
                        fabricante = :fabricante,
                        fabricante_norm = :fabricante_norm,
                        tamano = COALESCE(:tamano, tamano),
                        version = :version,
                        raw_line = :raw_line
                    WHERE id = :row_id
                    """
                ),
                {**payload, "raw_line": raw_line, "row_id": existing["id"]},
            )
        else:
            db.execute(
                text(
                    """
                    INSERT INTO software_instalado (
                        importacion_id, equipo_id, departamento_id, periodo, source_row,
                        nombre, nombre_norm, fabricante, fabricante_norm,
                        fecha_instalacion, tamano, version, raw_line
                    )
                    VALUES (
                        :importacion_id, :equipo_id, :departamento_id, :periodo, :source_row,
                        :nombre, :nombre_norm, :fabricante, :fabricante_norm,
                        NULL, :tamano, :version, :raw_line
                    )
                    """
                ),
                {
                    **payload,
                    "importacion_id": importacion_id,
                    "equipo_id": equipo_id,
                    "departamento_id": departamento_id,
                    "periodo": periodo,
                    "source_row": _next_source_row(db, importacion_id),
                    "raw_line": raw_line,
                },
            )

    for importacion_id in sorted(touched_imports):
        _refresh_program_count(db, importacion_id)


def dashboard_simple(periodo: str, db=None) -> dict:
    if db is None:
        with _engine().connect() as conn:
            return dashboard_simple(periodo, db=conn)
    periodo = validate_periodo(periodo)
    row = db.execute(
        text(
            """
            SELECT
                (SELECT COUNT(*) FROM equipos WHERE activo = TRUE) AS equipos,
                (SELECT COUNT(*) FROM departamentos) AS departamentos,
                (SELECT COUNT(*) FROM software_importaciones WHERE periodo = :periodo AND estado = 'confirmed') AS importaciones,
                (
                    SELECT COUNT(DISTINCT nombre_norm)
                    FROM software_instalado si
                    JOIN software_importaciones imp ON imp.id = si.importacion_id
                    WHERE imp.periodo = :periodo
                      AND imp.estado = 'confirmed'
                ) AS software
            """
        ),
        {"periodo": periodo},
    ).mappings().first()
    return dict(row) if row else {}


def resumen_departamentos(periodo: str, db=None) -> list[dict]:
    if db is None:
        with _engine().connect() as conn:
            return resumen_departamentos(periodo, db=conn)
    periodo = validate_periodo(periodo)
    rows = db.execute(
        text(
            """
            SELECT
                d.id,
                d.nombre,
                COALESCE(eq.total, 0) AS equipos,
                COALESCE(imp.total, 0) AS importados,
                COALESCE(sw.total, 0) AS software
            FROM departamentos d
            LEFT JOIN (
                SELECT departamento_id, COUNT(*) AS total
                FROM equipos
                WHERE activo = TRUE
                GROUP BY departamento_id
            ) eq ON eq.departamento_id = d.id
            LEFT JOIN (
                SELECT departamento_id, COUNT(*) AS total
                FROM software_importaciones
                WHERE periodo = :periodo AND estado = 'confirmed'
                GROUP BY departamento_id
            ) imp ON imp.departamento_id = d.id
            LEFT JOIN (
                SELECT imp.departamento_id, COUNT(DISTINCT si.nombre_norm) AS total
                FROM software_instalado si
                JOIN software_importaciones imp ON imp.id = si.importacion_id
                WHERE imp.periodo = :periodo AND imp.estado = 'confirmed'
                GROUP BY imp.departamento_id
            ) sw ON sw.departamento_id = d.id
            ORDER BY d.id
            """
        ),
        {"periodo": periodo},
    ).mappings().all()
    result = []
    for row in rows:
        item = dict(row)
        item["estado"] = f"{item['importados']} / {item['equipos']}"
        result.append(item)
    return result


def _cell_value(value: Any):
    if isinstance(value, bool):
        return "Si" if value else "No"
    return value


def _safe_sheet_name(name: str) -> str:
    invalid = "[]:*?/\\"
    cleaned = "".join("_" if c in invalid else c for c in str(name or "Hoja"))
    return cleaned[:31] or "Hoja"


def _write_rows(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="366092")
    for row in rows:
        ws.append(row)
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 70)
        ws.column_dimensions[column[0].column_letter].width = width
    ws.freeze_panes = "A2"


def exportar_inventario_excel(periodo: str, db=None) -> bytes:
    if db is None:
        with _engine().connect() as conn:
            return exportar_inventario_excel(periodo, db=conn)
    from modules.equipos import listar_equipos
    from modules.software import listar_departamentos

    periodo = validate_periodo(periodo)
    wb = Workbook()
    wb.remove(wb.active)

    equipos = listar_equipos(db, solo_activos=False)
    ws_eq = wb.create_sheet("Equipos")
    _write_rows(
        ws_eq,
        ["Departamento", "Equipo", "Usuario", "Activo"],
        [[e.get("departamento_nombre"), e.get("nombre"), e.get("notas"), _cell_value(e.get("activo"))] for e in equipos],
    )

    for dept in listar_departamentos(db):
        rows = software_por_departamento(periodo, dept["id"], db=db)
        ws = wb.create_sheet(_safe_sheet_name(dept["nombre"]))
        _write_rows(
            ws,
            ["Programa", "Editor/Fabricante", "Versiones", "N equipos", "Equipos"],
            [
                [
                    row.get("nombre"),
                    row.get("fabricantes"),
                    row.get("versiones"),
                    row.get("n_equipos"),
                    row.get("equipos"),
                ]
                for row in rows
            ],
        )

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
