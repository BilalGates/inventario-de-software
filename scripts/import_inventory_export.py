from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import text


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from database.connection import get_engine
from modules.equipos import importar_equipos_desde_lista
from modules.simple_inventory import (
    _program_payload,
    confirm_software_import,
    software_por_departamento,
    validate_periodo,
)
from modules.software import listar_departamentos
from utils.normalizer import normalize_equipo_nombre, normalize_nombre


EQUIPMENT_SHEET = "Equipos"
EQUIPMENT_COLUMNS = {
    "Departamento": "departamento_nombre",
    "Equipo": "nombre",
    "Usuario": "notas",
    "Activo": "activo",
    "Servidor": "es_servidor",
    "Tipo": "tipo_dispositivo",
    "Marca/Modelo": "marca_modelo",
    "N serie": "num_serie",
    "MAC": "mac_address",
    "Sistema operativo": "sistema_operativo",
    "Procesador": "procesador",
    "RAM": "ram",
    "Almacenamiento": "almacenamiento",
    "Responsable": "responsable",
    "Ubicacion": "ubicacion",
    "Coste": "coste",
    "Fecha adquisicion": "fecha_adquisicion",
    "Fecha alta": "fecha_alta",
    "Fecha baja": "fecha_baja",
}
SOFTWARE_HEADERS = ["Programa", "Editor/Fabricante", "Versiones", "N equipos", "Equipos"]


def _key(value: Any) -> str:
    raw = unicodedata.normalize("NFKD", str(value or "").strip().casefold())
    return " ".join("".join(ch for ch in raw if not unicodedata.combining(ch)).split())


def _safe_sheet_name(name: str) -> str:
    invalid = "[]:*?/\\"
    return "".join("_" if ch in invalid else ch for ch in name)[:31]


def _bool(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    normalized = _key(value)
    if normalized in {"si", "s", "yes", "true", "1", "activo", "servidor"}:
        return True
    if normalized in {"no", "n", "false", "0", "inactivo"}:
        return False
    return default


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _infer_period(path: Path) -> str | None:
    match = re.search(r"(?<!\d)(20\d{2}-(?:0[1-9]|1[0-2]))(?!\d)", path.stem)
    return match.group(1) if match else None


def _read_equipment(sheet) -> list[dict]:
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows, ()))
    mapping = {idx: EQUIPMENT_COLUMNS[value] for idx, value in enumerate(headers) if value in EQUIPMENT_COLUMNS}
    required = {"departamento_nombre", "nombre", "activo"}
    if not required.issubset(mapping.values()):
        raise ValueError(f"Cabeceras no reconocidas en '{EQUIPMENT_SHEET}': {headers}")

    result: list[dict] = []
    for row_number, row in enumerate(rows, start=2):
        item = {field: row[idx] for idx, field in mapping.items() if idx < len(row)}
        if not str(item.get("nombre") or "").strip():
            continue
        item["nombre"] = str(item["nombre"]).strip()
        item["departamento_nombre"] = str(item.get("departamento_nombre") or "").strip()
        item["activo"] = _bool(item.get("activo"), default=True)
        item["es_servidor"] = _bool(item.get("es_servidor"))
        item["fecha_adquisicion"] = _date(item.get("fecha_adquisicion"))
        item["fecha_alta"] = _date(item.get("fecha_alta"))
        item["fecha_baja"] = _date(item.get("fecha_baja"))
        item["_row"] = row_number
        result.append(item)
    return result


def _split_devices(value: Any) -> list[str]:
    return [part.strip() for part in str(value or "").split(",") if part.strip()]


def _read_software(workbook, departments: list[dict]) -> tuple[dict[tuple[str, str], list[dict]], dict]:
    by_device: dict[tuple[str, str], list[dict]] = defaultdict(list)
    stats = {
        "aggregate_rows": 0,
        "links": 0,
        "multi_version_rows": 0,
        "multi_publisher_rows": 0,
        "distributed_version_rows": 0,
        "distributed_publisher_rows": 0,
        "expanded_program_rows": 0,
    }
    sheet_names = {_key(name): name for name in workbook.sheetnames if name != EQUIPMENT_SHEET}

    for department in departments:
        expected = _safe_sheet_name(department["nombre"])
        sheet_name = sheet_names.get(_key(expected))
        if not sheet_name:
            raise ValueError(f"Falta la hoja del departamento '{department['nombre']}' (esperada: '{expected}').")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        if headers[:5] != SOFTWARE_HEADERS:
            raise ValueError(f"Cabeceras no reconocidas en '{sheet_name}': {headers}")

        department_key = _key(department["nombre"])
        for row_number, row in enumerate(rows, start=2):
            name = str(row[0] or "").strip()
            if not name:
                continue
            publisher = str(row[1] or "").strip() or None
            version = str(row[2] or "").strip() or None
            devices = _split_devices(row[4])
            expected_count = int(row[3] or 0)
            if expected_count != len(devices):
                raise ValueError(
                    f"{sheet_name} fila {row_number}: N equipos={expected_count}, pero la lista contiene {len(devices)}."
                )
            if not devices:
                raise ValueError(f"{sheet_name} fila {row_number}: el programa no tiene equipos asociados.")
            if version and ", " in version:
                stats["multi_version_rows"] += 1
            if publisher and ", " in publisher:
                stats["multi_publisher_rows"] += 1
            versions = [version]
            publishers = [publisher]
            if version and len(version) > 200:
                versions = [part.strip() for part in version.split(", ") if part.strip()]
                if any(len(part) > 200 for part in versions):
                    raise ValueError(
                        f"{sheet_name} fila {row_number}: una version individual supera los 200 caracteres."
                    )
                stats["distributed_version_rows"] += 1
            if publisher and len(publisher) > 300:
                publishers = [part.strip() for part in publisher.split(", ") if part.strip()]
                if any(len(part) > 300 for part in publishers):
                    raise ValueError(
                        f"{sheet_name} fila {row_number}: un fabricante individual supera los 300 caracteres."
                    )
                stats["distributed_publisher_rows"] += 1
            stats["aggregate_rows"] += 1
            stats["links"] += len(devices)
            assignment_count = max(len(devices), len(versions), len(publishers))
            stats["expanded_program_rows"] += assignment_count
            for assignment_index in range(assignment_count):
                device = devices[assignment_index % len(devices)]
                device_version = versions[assignment_index % len(versions)]
                device_publisher = publishers[assignment_index % len(publishers)]
                payload = _program_payload(
                    {"nombre": name, "fabricante": device_publisher, "version": device_version}
                )
                payload.update({"row_number": len(by_device[(department_key, _key(device))]) + 1, "raw_line": name})
                by_device[(department_key, _key(device))].append(payload)
    return by_device, stats


def _raw_text(rows: list[dict]) -> str:
    return "\n".join(
        "\t".join(
            [
                str(row.get("nombre") or ""),
                str(row.get("fabricante") or ""),
                "",
                "",
                str(row.get("version") or ""),
            ]
        )
        for row in rows
    )


def _write_snapshot(db, period: str, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    destination = backup_dir / f"pre_import_inventory_{period}_{stamp}.json"
    snapshot = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "period": period,
        "equipos": [dict(row) for row in db.execute(text("SELECT * FROM equipos ORDER BY id")).mappings()],
        "software_importaciones": [
            dict(row)
            for row in db.execute(
                text("SELECT * FROM software_importaciones WHERE periodo = :period ORDER BY id"),
                {"period": period},
            ).mappings()
        ],
        "software_instalado": [
            dict(row)
            for row in db.execute(
                text("SELECT * FROM software_instalado WHERE periodo = :period ORDER BY id"),
                {"period": period},
            ).mappings()
        ],
    }
    destination.write_text(json.dumps(snapshot, ensure_ascii=False, default=str, indent=2), encoding="utf-8")
    return destination


def import_export(path: Path, period: str, apply: bool, backup_dir: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if EQUIPMENT_SHEET not in workbook.sheetnames:
        raise ValueError(f"Falta la hoja obligatoria '{EQUIPMENT_SHEET}'.")

    engine = get_engine()
    with engine.begin() as db:
        departments = listar_departamentos(db)
        equipment = _read_equipment(workbook[EQUIPMENT_SHEET])
        programs_by_device, software_stats = _read_software(workbook, departments)

        department_keys = {_key(d["nombre"]): d for d in departments}
        source_devices: dict[tuple[str, str], dict] = {}
        errors: list[str] = []
        for item in equipment:
            department_key = _key(item["departamento_nombre"])
            device_key = (department_key, _key(item["nombre"]))
            if department_key not in department_keys:
                errors.append(
                    f"Equipos fila {item['_row']}: departamento desconocido '{item['departamento_nombre']}'."
                )
            if device_key in source_devices:
                errors.append(f"Equipo duplicado en el Excel: {item['departamento_nombre']} / {item['nombre']}.")
            source_devices[device_key] = item

        for device_key in programs_by_device:
            item = source_devices.get(device_key)
            if not item:
                errors.append(f"El software referencia un equipo que no figura en Equipos: {device_key}.")
            elif not item["activo"]:
                errors.append(f"El software referencia un equipo inactivo: {item['nombre']}.")
        if errors:
            raise ValueError("\n".join(errors[:30]))

        existing_confirmed = db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM software_importaciones
                WHERE periodo = :periodo AND estado = 'confirmed'
                """
            ),
            {"periodo": period},
        ).scalar_one()

        summary = {
            "file": str(path),
            "period": period,
            "mode": "apply" if apply else "dry-run",
            "departments": len(departments),
            "equipment_rows": len(equipment),
            "active_equipment": sum(1 for item in equipment if item["activo"]),
            "devices_with_software": len(programs_by_device),
            "existing_confirmed_imports_for_period": int(existing_confirmed),
            **software_stats,
        }
        if not apply:
            db.rollback()
            return summary

        summary["backup"] = str(_write_snapshot(db, period, backup_dir))

        clean_equipment = [{k: v for k, v in item.items() if not k.startswith("_")} for item in equipment]
        equipment_result = importar_equipos_desde_lista(db, clean_equipment, departments)
        if equipment_result["errores"]:
            raise ValueError("\n".join(equipment_result["errores"]))

        device_ids: dict[tuple[str, str], int] = {}
        for item in equipment:
            department = department_keys[_key(item["departamento_nombre"])]
            device = db.execute(
                text(
                    """
                    SELECT id FROM equipos
                    WHERE departamento_id = :department_id AND nombre_norm = :name_norm
                    """
                ),
                {
                    "department_id": department["id"],
                    "name_norm": normalize_equipo_nombre(item["nombre"]),
                },
            ).scalar_one()
            device_ids[(_key(item["departamento_nombre"]), _key(item["nombre"]))] = int(device)
            db.execute(
                text(
                    """
                    UPDATE equipos
                    SET fecha_alta = COALESCE(:fecha_alta, fecha_alta),
                        fecha_baja = :fecha_baja
                    WHERE id = :id
                    """
                ),
                {"id": device, "fecha_alta": item.get("fecha_alta"), "fecha_baja": item.get("fecha_baja")},
            )

        imported_programs = 0
        imported_devices = 0
        for device_key, rows in programs_by_device.items():
            confirm_software_import(device_ids[device_key], period, rows, _raw_text(rows), db=db)
            imported_devices += 1
            imported_programs += len(rows)

        summary["equipment_result"] = equipment_result
        summary["imported_devices"] = imported_devices
        summary["imported_program_rows"] = imported_programs
        return summary


def verify_import(path: Path, period: str, backup_dir: Path) -> dict:
    expected = import_export(path, period, False, backup_dir)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    aggregate_mismatches: list[str] = []
    with get_engine().connect() as db:
        actual = dict(
            db.execute(
                text(
                    """
                    SELECT
                        (SELECT COUNT(*) FROM equipos) AS equipment,
                        (SELECT COUNT(*) FROM equipos WHERE activo = TRUE) AS active_equipment,
                        (
                            SELECT COUNT(*) FROM software_importaciones
                            WHERE periodo = :period AND estado = 'confirmed'
                        ) AS confirmed_imports,
                        (
                            SELECT COUNT(*) FROM software_instalado
                            WHERE periodo = :period
                        ) AS installed_rows,
                        (
                            SELECT COUNT(*) FROM (
                                SELECT departamento_id, nombre_norm
                                FROM software_instalado
                                WHERE periodo = :period
                                GROUP BY departamento_id, nombre_norm
                            ) grouped_programs
                        ) AS aggregate_rows,
                        (
                            SELECT COUNT(*) FROM (
                                SELECT equipo_id
                                FROM software_importaciones
                                WHERE periodo = :period AND estado = 'confirmed'
                                GROUP BY equipo_id
                                HAVING COUNT(*) > 1
                            ) duplicate_imports
                        ) AS duplicate_confirmed_imports
                    """
                ),
                {"period": period},
            ).mappings().one()
        )
        departments = listar_departamentos(db)
        sheet_names = {_key(name): name for name in workbook.sheetnames if name != EQUIPMENT_SHEET}
        for department in departments:
            sheet_name = sheet_names[_key(_safe_sheet_name(department["nombre"]))]
            source_rows = list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))
            source = {
                normalize_nombre(row[0]): (
                    str(row[0] or "").strip(),
                    str(row[1] or "").strip() or None,
                    str(row[2] or "").strip() or None,
                    int(row[3] or 0),
                    str(row[4] or "").strip() or None,
                )
                for row in source_rows
                if str(row[0] or "").strip()
            }
            saved = {
                row["nombre_norm"]: (
                    str(row["nombre"] or "").strip(),
                    row["fabricantes"],
                    row["versiones"],
                    int(row["n_equipos"]),
                    row["equipos"],
                )
                for row in software_por_departamento(period, department["id"], db=db)
            }
            if source != saved:
                missing = sorted(set(source) - set(saved))
                extra = sorted(set(saved) - set(source))
                changed = sorted(key for key in set(source) & set(saved) if source[key] != saved[key])
                aggregate_mismatches.append(
                    f"{department['nombre']}: missing={len(missing)}, extra={len(extra)}, changed={len(changed)}"
                )
    checks = {
        "confirmed_imports": actual["confirmed_imports"] == expected["devices_with_software"],
        "installed_rows": actual["installed_rows"] == expected["expanded_program_rows"],
        "aggregate_rows": actual["aggregate_rows"] == expected["aggregate_rows"],
        "no_duplicate_confirmed_imports": actual["duplicate_confirmed_imports"] == 0,
        "aggregate_values_match": not aggregate_mismatches,
    }
    return {
        "period": period,
        "expected": expected,
        "actual": actual,
        "aggregate_mismatches": aggregate_mismatches,
        "checks": checks,
        "ok": all(checks.values()),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Importa una exportacion completa de Inventario Asserta.")
    parser.add_argument("--file", required=True, type=Path, help="Ruta al Excel exportado por Inventario Asserta.")
    parser.add_argument("--period", help="Periodo YYYY-MM; si se omite, se toma del nombre del fichero.")
    parser.add_argument("--apply", action="store_true", help="Confirma la escritura. Sin esta opcion solo valida.")
    parser.add_argument("--verify-only", action="store_true", help="Compara el Excel con los conteos guardados.")
    parser.add_argument("--backup-dir", type=Path, default=ROOT / "backups", help="Directorio para el snapshot previo.")
    args = parser.parse_args()

    path = args.file.expanduser().resolve()
    if not path.is_file():
        parser.error(f"No existe el fichero: {path}")
    period = validate_periodo(args.period or _infer_period(path) or "")
    backup_dir = args.backup_dir.expanduser().resolve()
    if args.verify_only:
        result = verify_import(path, period, backup_dir)
    else:
        result = import_export(path, period, args.apply, backup_dir)
    print(json.dumps(result, ensure_ascii=False, default=str, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
